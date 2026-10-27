from pathlib import Path
from datetime import datetime, timedelta
import csv
import re
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / 'india_dataset.xlsx'
OUT = ROOT / 'processed'
OUT.mkdir(exist_ok=True)

wb = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
ws = wb.active
rows = ws.iter_rows(values_only=True)
for _ in range(5):
    next(rows)
header = next(rows)
dates = [pd.to_datetime(v).normalize() for v in header[1:] if v is not None]
if len(dates) != ws.max_column - 1:
    raise ValueError('Unexpected empty date headers')

records = []
for row in rows:
    label = str(row[0] or '').strip()
    match = re.match(r'^(\d{1,2}):(\d{2})\s*-', label)
    if not match:
        continue
    hour, minute = map(int, match.groups())
    if hour > 23 or minute not in (0, 15, 30, 45):
        raise ValueError(f'Invalid time block: {label}')
    for day, value in zip(dates, row[1:]):
        power = pd.to_numeric(value, errors='coerce')
        records.append((day + pd.Timedelta(hours=hour, minutes=minute), power))

pv = pd.DataFrame(records, columns=['Timestamp', 'PV_Power_MW'])
pv = pv.sort_values('Timestamp').reset_index(drop=True)
if pv.Timestamp.duplicated().any():
    raise ValueError('Duplicate timestamps found')
expected = pd.date_range(pv.Timestamp.min(), pv.Timestamp.max(), freq='15min')
if len(expected) != len(pv) or not pv.Timestamp.equals(pd.Series(expected, name='Timestamp')):
    raise ValueError('15-minute time sequence is incomplete')
if (pv.PV_Power_MW.dropna() < 0).any():
    raise ValueError('Negative generation found: inspect before processing')
if (pv.PV_Power_MW.dropna() > 5).any():
    raise ValueError('Generation above 5 MW plant capacity: inspect')

pv.to_csv(OUT / 'india_pv_15min.csv', index=False)
hourly = pv.set_index('Timestamp').resample('h').agg(
    PV_Power_MW=('PV_Power_MW', 'mean'),
    Valid_15min_Readings=('PV_Power_MW', 'count')
)
hourly.loc[hourly.Valid_15min_Readings < 4, 'PV_Power_MW'] = float('nan')
hourly.reset_index().to_csv(OUT / 'india_pv_hourly.csv', index=False)

daily = pv.assign(Date=pv.Timestamp.dt.date).groupby('Date').agg(
    Available_Readings=('PV_Power_MW', 'count'),
    Missing_Readings=('PV_Power_MW', lambda s: int(s.isna().sum()))
).reset_index()
daily.to_csv(OUT / 'india_daily_quality.csv', index=False)

print('Source:', SOURCE.name)
print('Dates:', dates[0].date(), 'to', dates[-1].date())
print('15-minute records:', len(pv))
print('Missing readings:', int(pv.PV_Power_MW.isna().sum()))
print('Days with missing readings:', int((daily.Missing_Readings > 0).sum()))
print('Completely missing days:', int((daily.Available_Readings == 0).sum()))
print('Hourly records:', len(hourly))
print('Valid complete hours:', int(hourly.PV_Power_MW.notna().sum()))
print('Outputs:', OUT)
